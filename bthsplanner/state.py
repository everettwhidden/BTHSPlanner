
import io
import math
import re
import openpyxl
import reflex as rx
from data.courses import double_alternates
from data.courses import alternate_list
from data.courses import COURSE_CATALOG, MAJORS, req_courses
from data.courses import alternates_with_gym

MAJOR_NAMES = list(MAJORS.keys())


def _next_lang_level(lang: str) -> str:
    if lang.endswith(" I"):
        return lang[:-1] + "II"
    return lang  # already at level II, nothing higher to bump to

def _grade_badges(is_elective: bool, grades: list, available_grades: list) -> list[str]:
    if grades:
        return grades
    elif available_grades:
        return available_grades
    elif is_elective:
        return ["Junior", "Senior"]
    else:
        return []

def _build_partner_map(pairs):
    m = {}
    for a, b in pairs:
        m[a] = b
        m[b] = a
    return m


    
class PlannerState(rx.State):
    
    
    active_tab: str = "Plan"
    course_catalog: list[dict] = COURSE_CATALOG
    modal_source: str = "planner"
    is_detail_open: bool = False

    pre_credits: dict[str, bool] = {
        "Algebra I": False,
        "Geometry": False,
        "US History" : False,
        "Biology/Living Environment" : False,
    }
    

    partner_map: dict[str, str] = _build_partner_map(double_alternates)


    def toggle_pre_credit(self, credit: str, checked: bool):
        snapshot = self._export_progress_dict()
        self.pre_credits[credit] = checked
        if credit == "Geometry":
            if checked:
        
                self.j_math = "AP Calculus BC"
                self.sn_math = "Multivariable Calculus"
            else:
                self.j_math = "Pre-Calculus"
                self.sn_math = "AP Calculus BC"
        if credit == "US History":
            if checked:
                self.j_his = "Economics"
            elif self.j_his == "Economics":
                self.j_his = "Regents United States History"
            else:
                self.j_his = "Regents United States History"
        if credit == "Language":
            if checked:
                self.f_lang = _next_lang_level(self.pre_language)
            else:
                self.f_lang = self.pre_language
            self.sync_so_eng_for_language()
        if credit in ("Biology/Living Environment", "Algebra I"):
            self.sync_f_sci()
        self._enforce_ap_limits()
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)

    pre_language: str = "Spanish I"

    def change_pre_language(self, val: str):
        self.pre_language = val
        if self.pre_credits.get("Language", False):
            self.f_lang = _next_lang_level(val)
        else:
            self.f_lang = val

    
    def sync_f_sci(self):
        has_bio = self.pre_credits.get("Biology/Living Environment", False)
        
        if not has_bio:
            self.f_sci = "Regents Biology"
        else:
            self.f_sci = "Regents Physics"

    def sync_so_eng_for_language(self):
        if self.sophomore_ap_count > 3 and self.so_eng == "AP Capstone Seminar":
            self.so_eng = "English 10"

    

    def change_tab(self, name: str):
        self.active_tab = name
    #initial courses and major
    major: str = "Aerospace Engineering"
    f_hist: str = "Pre-AP World History"
    f_sci: str = "Regents Physics"
    f_math: str = "Geometry"
    f_lang: str = "Spanish I"
    so_eng: str = "English 10"
    so_hist: str = "AP European History"
    j_eng: str = "English 11"
    j_his: str = "Regents United States History"
    j_phys: str = "Regents Physics"
    j_math: str = "Pre-Calculus"
    sn_math: str = "AP Calculus BC"
    sn_math_active: bool = True
    sn_eng: str = "Creative Writing"
    sn_his: str = "Economics"
    sn_ec: str = "Economics"
    sn_gov: str = "Participation in Government"

    # elective tracking lists 
    freshman_electives: list[str] = []
    sophomore_electives: list[str] = []
    junior_electives: list[str] = []   
    senior_electives: list[str] = []   
    

    
        
    
    # import/export data functions

    def _export_progress_dict(self) -> dict:
        return {
            "major": self.major,
            "f_hist": self.f_hist, "f_sci": self.f_sci, "f_math": self.f_math, "f_lang": self.f_lang,
            "so_eng": self.so_eng, "so_hist": self.so_hist,
            "j_eng": self.j_eng, "j_his": self.j_his, "j_phys": self.j_phys, "j_math": self.j_math,
            "sn_eng": self.sn_eng, "sn_his": self.sn_his, "sn_ec": self.sn_ec, "sn_gov": self.sn_gov, "sn_math": self.sn_math,
            "freshman_electives": list(self.freshman_electives),
            "sophomore_electives": list(self.sophomore_electives),
            "junior_electives": list(self.junior_electives),
            "senior_electives": list(self.senior_electives),
        }

    def _import_progress_dict(self, data: dict):
        self.major = data.get("major", self.major)
        self.f_hist = data.get("f_hist", self.f_hist)
        self.f_sci = data.get("f_sci", self.f_sci)
        self.f_math = data.get("f_math", self.f_math)
        self.f_lang = data.get("f_lang", self.f_lang)
        self.so_eng = data.get("so_eng", self.so_eng)
        self.so_hist = data.get("so_hist", self.so_hist)
        self.j_eng = data.get("j_eng", self.j_eng)
        self.j_his = data.get("j_his", self.j_his)
        self.j_phys = data.get("j_phys", self.j_phys)
        self.j_math = data.get("j_math", self.j_math)
        self.sn_eng = data.get("sn_eng", self.sn_eng)
        self.sn_his = data.get("sn_his", self.sn_his)
        self.sn_ec = data.get("sn_ec", self.sn_ec)
        self.sn_gov = data.get("sn_gov", self.sn_gov)
        self.sn_math = data.get("sn_math", self.sn_math)
        self.freshman_electives = data.get("freshman_electives", [])
        self.sophomore_electives = data.get("sophomore_electives", [])
        self.junior_electives = data.get("junior_electives", [])
        self.senior_electives = data.get("senior_electives", [])


    #search and filter stuff

    search_query: str = ""             # Live search text string


    explore_search_query: str = ""

    explore_search_query: str = ""

    order: str = "Default"
    selected_grades: list[str] = []
    selected_majors: list[str] = []

    def set_explore_search_query(self, val: str):
        self.explore_search_query = val

    def set_order(self, val: str):
        self.order = val

    def toggle_grade(self, grade: str, checked: bool):
        if checked and grade not in self.selected_grades:
            self.selected_grades.append(grade)
        elif not checked and grade in self.selected_grades:
            self.selected_grades.remove(grade)

    def toggle_major(self, major_name: str, checked: bool):
        if checked and major_name not in self.selected_majors:
            self.selected_majors.append(major_name)
        elif not checked and major_name in self.selected_majors:
            self.selected_majors.remove(major_name)

    def compute_explore_catalog(self) -> list[dict]:
        query = self.explore_search_query.lower().strip()
        results = []
        for name, data in COURSE_CATALOG.items():
            if query not in name.lower() and query not in data.get("description", "").lower():
                continue

            badges = _grade_badges(
                data.get("is_elective", False),
                data.get("grades", []),
                data.get("available_grades", [])
            )

            # grade filter, keep the course if it matches any picked grade
            if self.selected_grades:
                if not any(g in badges for g in self.selected_grades):
                    continue

            # major filter, keep the course if it's in any picked major
            if self.selected_majors:
                in_major = False
                for m in self.selected_majors:
                    major_data = MAJORS.get(m, {})
                    if name in major_data.get("Junior", []) or name in major_data.get("Senior", []):
                        in_major = True
                        break
                if not in_major:
                    continue

            results.append({
                "name": name,
                "description": data.get("description", ""),
                "periods": data.get("periods", 1),
                "semesters": data.get("semesters", 2),
                "credits": f"{float(data.get('credits', 1.0)):.1f}",
                "is_elective": data.get("is_elective", False),
                "disabled": False,
                "grade_badges": badges,
            })

        if self.order == "A-Z":
            results = sorted(results, key=lambda c: c["name"].lower())

        return results
    @rx.var
    def explore_catalog(self) -> list[dict]:
        return self.compute_explore_catalog()

    @rx.var
    def explore_catalog_count(self) -> int:
        return len(self.compute_explore_catalog())
    


    active_modal_course: dict = {}     # Holds the course info for detail popup
    is_detail_open: bool = False       # Controls the visibility of the global detail modal

    #adnvaced lang courses electives, exclude spanish bc spanihs has advanced spanish
    ap_lotes = ["AP Chinese Language & Culture", "AP French Language & Culture", "AP German Language & Culture", "AP Italian Language & Culture"]
    none_hisory_placeholders = {
    "None (free period, econ senior year)",
    "None (free period, econ junior year)",
}

    current_modal_grade: str = "Junior"
    
    @rx.var
    def current_page(self) -> str:
        return self.router.page.path
    
    # --- Existing Explicit Mutator Actions ---
    def export_to_excel(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "4-Year Plan"

        years = ["Freshman", "Sophomore", "Junior", "Senior"]
        year_cols = {"Freshman": 2, "Sophomore": 3, "Junior": 4, "Senior": 5}  # B, C, D, E

        # --- Row 1: Major ---
        ws.cell(row=1, column=1, value="Major")
        ws.cell(row=1, column=2, value=self.major)

        # --- Row 3: Year headers ---
        ws.cell(row=3, column=1, value="Year")
        for year, col in year_cols.items():
            ws.cell(row=3, column=col, value=year)

        # --- Rows 4+: courses, one per row, stacked under each year's column ---
        year_courses = {year: sorted(self._courses_for_year(year)) for year in years}
        max_courses = max(len(c) for c in year_courses.values())

        for i in range(max_courses):
            row = 4 + i
            for year, col in year_cols.items():
                courses = year_courses[year]
                if i < len(courses):
                    ws.cell(row=row, column=col, value=courses[i])

        # lean rows that summarize the totals for eachb category
        summary_start = 4 + max_courses + 1
        ws.cell(row=summary_start, column=1, value="Courses")
        ws.cell(row=summary_start + 1, column=1, value="Periods")
        ws.cell(row=summary_start + 2, column=1, value="AP Count")

        course_counts = {
            "Freshman": self.freshman_course_count, "Sophomore": self.sophomore_course_count,
            "Junior": self.junior_course_count, "Senior": self.senior_course_count,
        }
        total_courses = sum(course_counts.values())
        periods = {
            "Freshman": self.freshman_periods, "Sophomore": self.sophomore_periods,
            "Junior": self.junior_periods, "Senior": self.senior_periods,
        }
        total_periods = sum(periods.values())
        ap_counts = {
            "Freshman": self.freshman_ap_count, "Sophomore": self.sophomore_ap_count,
            "Junior": self.junior_ap_count, "Senior": self.senior_ap_count,
        }
        total_ap = sum(ap_counts.values())

        for year, col in year_cols.items():
            ws.cell(row=summary_start, column=col, value=course_counts[year])
            ws.cell(row=summary_start, column=col+1, value=total_courses)
            ws.cell(row=summary_start + 1, column=col, value=periods[year])
            ws.cell(row=summary_start + 1, column=col+1, value=total_periods)
            ws.cell(row=summary_start + 2, column=col, value=ap_counts[year])
            ws.cell(row=summary_start + 2, column=col+1, value=total_ap)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return rx.download(
            data=buffer.getvalue(),
            filename="brooklyn_tech_plan.xlsx",
        )

    @rx.var
    def freshman_math_courses(self) -> list[str]:
        if self.math_is_accelerated:
            return ["Algebra II"]
        if not self.pre_credits.get("Algebra I", False):
            return ["Algebra I", "Geometry"]
        return ["Geometry"]
    
    def change_major(self, val: str):
        self.major = val
        # Reset all dropdowns and electives to defaults on major change
        self.f_hist = "Pre-AP World History"
        self.sync_f_sci()
        self.f_math = "Geometry"
        if self.pre_credits.get("Language", False):
            self.f_lang = _next_lang_level(self.pre_language)
        else:
            self.f_lang = self.pre_language
        self.so_eng = "English 10"
        self.so_hist = "AP European History"
        self.j_eng = "English 11"
        self.j_his = "Regents United States History"
        self.j_phys = "Regents Physics"
        self.j_math = self.j_math if self.math_is_accelerated else "Pre-Calculus"
        self.sn_eng = "Creative Writing"
        self.sn_his = "Participation in Government" if self.has_econ_gov_conflict else "Economics"
        self.sn_ec = "Economics"
        self.sn_gov = "Participation in Government"
        self.freshman_electives = []
        self.sophomore_electives = []
        self.junior_electives = []
        self.senior_electives = []
        self.sn_math = "Multivariable Calculus" if self.math_is_accelerated else "AP Calculus BC"


    def change_f_hist(self, val: str):
        snapshot = self._export_progress_dict()
        self.f_hist = val
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)

    def change_f_sci(self, val: str):
        snapshot = self._export_progress_dict()
        self.f_sci = val
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)
    def change_f_math(self, val: str):
        snapshot = self._export_progress_dict()
        self.f_math = val
        self._enforce_ap_limits()
        self._evict_duplicates()
        self.j_math = "AP Calculus BC"
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)
    def change_f_lang(self, val: str):
        snapshot = self._export_progress_dict()
        self.f_lang = val
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)
    def change_so_eng(self, val: str):
        snapshot = self._export_progress_dict()
        self.so_eng = val
        if val == "AP Capstone Seminar":
    
            if self.senior_ap_count < 3:
                self.sn_eng = "AP Capstone Research"
                if not self._enforce_period_limits():
                    self._import_progress_dict(snapshot)
        else:
            self.sn_eng = "Creative Writing"
            if not self._enforce_period_limits():
                self._import_progress_dict(snapshot)

    def change_so_hist(self, val: str):
        snapshot = self._export_progress_dict()
        self.so_hist = val
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)

    def change_j_his(self, val: str):
        snapshot = self._export_progress_dict()
        prev = self.j_his
        self.j_his = val
        prev_needed_senior_econ = prev in ("Economics", "None (free period, econ senior year)")
        if val == "None (free period, econ senior year)" or (prev_needed_senior_econ and val != "Economics"):
            self.sn_his = "Economics"
            if "AP Macroeconomics with Government" in self.senior_electives:
                self.senior_electives.remove("AP Macroeconomics with Government")
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)
            if not self._enforce_period_limits():
                self._import_progress_dict(snapshot)

    def change_j_phys(self, val: str):
        snapshot = self._export_progress_dict()
        self.j_phys = val
        self._enforce_ap_limits()
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)

    def change_j_math(self, val: str):
        snapshot = self._export_progress_dict()
        self.j_math = val
        if val == "AP Calculus BC":
            self.sn_math = "Multivariable Calculus"
        if val == "Pre-Calculus":
            self.sn_math = "AP Calculus BC"
        self._enforce_ap_limits()
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)

    def change_sn_math(self, val: str):
        snapshot = self._export_progress_dict()
        self.sn_math = val
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)
    def toggle_sn_math(self):
        self.sn_math_active = not self.sn_math_active

    

    #for physcs majors physics II and or Physics c e and m
    selected_senior_physics: set[str] = {"AP Physics C: Electricity & Magnetism"}
    
        #selected_senior_physics.remove("AP Physics II")
    def toggle_senior_physics(self, course_name: str):
        #makes sure physics majors keep atleast one
        if course_name in self.selected_senior_physics:
            if len(self.selected_senior_physics) > 1:
                self.selected_senior_physics.remove(course_name)
        else:
            self.selected_senior_physics.add(course_name)
            if self.senior_periods > 10 or self.senior_ap_count > 3:
                self.selected_senior_physics.remove(course_name)
            

    def change_j_eng(self, val: str):
        snapshot = self._export_progress_dict()
        self.j_eng = val
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)
    def change_sn_eng(self, val: str):
        snapshot = self._export_progress_dict()
        self.sn_eng = val
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)
    def change_sn_his(self, val: str):
        snapshot = self._export_progress_dict()
        self.sn_his = val
        if val == "None (free period, econ junior year)" :
            self.j_his = "Economics"
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)
    def change_sn_ec(self, val: str):
        snapshot = self._export_progress_dict()
        self.sn_ec = val
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)
    def change_sn_gov(self, val: str):
        snapshot = self._export_progress_dict()
        self.sn_gov = val
        self._evict_duplicates()
        if not self._enforce_period_limits():
            self._import_progress_dict(snapshot)

    def _calculate_total_periods(self, year: str) -> int:
        courses = self._courses_for_year(year)
        full_year_periods = 0
        single_semester_1_pd = 0
        single_semester_2_pd = 0
        
        for c in courses:
            data = COURSE_CATALOG.get(c, {})
            pds = data.get("periods", 1)
            
            if data.get("semesters") == 1:
                if pds == 2:
                    # Count how many 2 period semester classes they have
                    single_semester_2_pd += 1
                else:
                    # Count how many 1 period semester classes they have
                    single_semester_1_pd += 1

            elif c in alternates_with_gym:
                full_year_periods += (pds-1)

            else:
                if pds == 2 and c in alternate_list:
                    full_year_periods += 1
                else:
                    full_year_periods += pds
                
        # 1. Pair up 1-period semester classes: math.ceil(count / 2) -> each pair = 1 period
        # 2. Pair up 2-period semester classes: math.ceil(count / 2) * 2 -> each pair = 2 periods
        total_semester_room = math.ceil(single_semester_1_pd / 2) + (math.ceil(single_semester_2_pd / 2) * 2)
        
        return full_year_periods + total_semester_room

    @rx.var
    def freshman_periods(self) -> int:
        return self._calculate_total_periods("Freshman")

    @rx.var
    def sophomore_periods(self) -> int:
        return self._calculate_total_periods("Sophomore")

    

    @rx.var
    def junior_periods(self) -> int:
        return self._calculate_total_periods("Junior")

    @rx.var
    def senior_periods(self) -> int:
        return self._calculate_total_periods("Senior")

    @rx.var
    def freshman_course_count(self) -> int:
        return len(self._courses_for_year("Freshman"))

    @rx.var
    def sophomore_course_count(self) -> int:
        return len(self._courses_for_year("Sophomore"))

    @rx.var
    def junior_course_count(self) -> int:
        return len(self._courses_for_year("Junior"))

    @rx.var
    def senior_course_count(self) -> int:
        return len(self._courses_for_year("Senior"))

    @rx.var
    def total_course_count(self) -> int:
        return (self.freshman_course_count + 
                self.sophomore_course_count + 
                self.junior_course_count + 
                self.senior_course_count)

    
    
    # --- AP Course Counters ---
    @rx.var
    def freshman_ap_count(self) -> int:
        count = 0
        for c in self._courses_for_year("Freshman"):
            # pre ap world is not ap
            if c[:2] == "AP":
                count += 1
        return count

    @rx.var
    def sophomore_ap_count(self) -> int:
        return sum(1 for c in self._courses_for_year("Sophomore") if "AP" in c)

    @rx.var
    def junior_ap_count(self) -> int:
        return sum(1 for c in self._courses_for_year("Junior") if "AP" in c)

    @rx.var
    def senior_ap_count(self) -> int:
        return sum(1 for c in self._courses_for_year("Senior") if "AP" in c)
    
  
    
    # --- Global Pop-up Event Actions ---
    def set_search_query(self, val: str):
        self.search_query = val

    def close_course_details(self):
        self.is_detail_open = False
    
    def open_course_details(self, name: str):
        data = COURSE_CATALOG.get(name, {})
        self.active_modal_course = {
            "name": name,
            "description": data.get("description", "No description available for this course yet."),
            "periods": str(data.get("periods", 1)),
            "credits": str(data.get("credits", 1.0)),
            "is_elective": data.get("is_elective", False),
            "semesters": str(data.get("semesters", 2)),
            "grade_badges": _grade_badges(
                data.get("is_elective", False),
                data.get("grades", []),
                data.get("available_grades", [])
            ),
        }
        self.is_detail_open = True

    def open_search_for_grade(self, grade_name: str):
        self.current_modal_grade = grade_name
        self.search_query = ""

    def add_course(self, course_name: str, target_year: str):
        if target_year == "Freshman" and course_name not in self.freshman_electives:
            self.freshman_electives.append(course_name)
        elif target_year == "Sophomore" and course_name not in self.sophomore_electives:
            self.sophomore_electives.append(course_name)
        elif target_year == "Junior" and course_name not in self.junior_electives:
            self.junior_electives.append(course_name)
        elif target_year == "Senior" and course_name not in self.senior_electives:
            self.senior_electives.append(course_name)
        self._evict_duplicates()

    def delete_course(self, course_name: str, target_year: str):
        junior_reqs = MAJORS[self.major].get("Junior", [])
        senior_reqs = MAJORS[self.major].get("Senior", [])
        
        if course_name in junior_reqs or course_name in senior_reqs:
            return rx.window_alert(f"Cannot delete {course_name}! This course is a core requirement for {self.major}.")
            
        if target_year == "Freshman" and course_name in self.freshman_electives:
            self.freshman_electives.remove(course_name)
        elif target_year == "Sophomore" and course_name in self.sophomore_electives:
            self.sophomore_electives.remove(course_name)
        elif target_year == "Junior" and course_name in self.junior_electives:
            self.junior_electives.remove(course_name)
        elif target_year == "Senior" and course_name in self.senior_electives:
            self.senior_electives.remove(course_name)

    # --- Computed Vars ---
    @rx.var(cache=True)
    def room(self) -> str: return MAJORS[self.major].get("room", "N/A")

    @rx.var(cache=True)
    def major_note(self) -> str: return MAJORS[self.major].get("note", "")

    @rx.var(cache=True)
    def math_is_accelerated(self) -> bool: return self.pre_credits.get("Geometry", False)

    @rx.var
    def took_ap_calc_bc_junior(self) -> bool:
        return self.math_is_accelerated or self.j_math == "AP Calculus BC" or self.j_math == "AP Calculus AB"

    @rx.var(cache=True)
    def lang_is_accelerated(self) -> bool: return "II" in self.f_lang

    @rx.var(cache=True)
    def language_base(self) -> str: return self.f_lang.split()[0] if self.f_lang else ""

    @rx.var(cache=True)
    def took_seminar(self) -> bool: return self.so_eng == "AP Capstone Seminar"

    @rx.var(cache=True)
    def chose_economics_jr(self) -> bool: 
        took_reg_econ = self.j_his == "Economics"
        return took_reg_econ
    
    @rx.var(cache=True)
    def chose_african_jr(self) -> bool:
        return self.j_his == "AP African American Studies" 
    @rx.var(cache=True)
    def chose_government_jr(self) -> bool:
        return (
            ("AP US Government" in self.junior_major_courses or 
            "AP US Government" in self.junior_electives) or ("AP US Government" in self.senior_major_courses or "AP US Government" in self.senior_electives) or ("AP Comparative Government" in self.senior_major_courses)
        )
    @rx.var(cache = True)
    def has_macro(self) -> bool:
        return ("AP Macroeconomics with Government" in self.junior_major_courses or "AP Macroeconomics with Government" in self.senior_major_courses)

    @rx.var(cache=True)
    def has_econ_gov_conflict(self) -> bool:
        #True if the student already has AP Macroeconomics or AP Macroeconomics / AP Microeconomics (One Semester each)
        conflict_courses = {
            "AP Macroeconomics with Government",
            "AP Macroeconomics / AP Microeconomics (One Semester each)",
        }
        sources = (
            self.junior_major_courses + self.senior_major_courses +
            self.junior_electives + self.senior_electives
        )
        return any(c in conflict_courses for c in sources)
    @rx.var(cache=True)
    def has_gov_conflict(self) -> bool:
        """True if the student's major already includes AP Comparative Government 
        as a required junior-year course (currently: LAS) — this conflicts with 
        also taking Participation in Government junior year, since Comparative 
        Government already satisfies the gov credit."""
        return self.major == "Law & Society" or "AP Comparative Government" in MAJORS[self.major].get("Junior", [])
    @rx.var(cache=True)

    def took_ap_hg_fr(self) -> bool:
        return "AP Human Geography" in self._courses_for_year("Freshman")
    
    @rx.var(cache=True)
    def has_ap_physics_junior(self) -> bool: return any("AP Physics I" in c for c in MAJORS[self.major].get("Junior", []))

    @rx.var(cache=True)
    def need_regents_physics_junior(self) -> bool: 
        if "AP Physics I" in self.junior_electives or "AP Physics I" in self.senior_electives:
            return False
        return (self.f_sci != "Regents Physics") and not self.has_ap_physics_junior

    @rx.var(cache=True)
    def junior_ap_room_for_physics(self) -> bool:
        """Whether there's room for another AP junior year, not counting the
        Regents Physics / AP Physics I swap slot itself."""
        others = [c for c in self._courses_for_year("Junior") if "AP" in c and c != self.j_phys]
        return len(others) < 3
    @rx.var(cache = True)
    def major_has_aphg(self) -> bool:
        return "AP Human Geography" in self.junior_major_courses or "AP Human Geography" in self.senior_major_courses

    @rx.var(cache=True)
    def junior_major_courses(self) -> list[str]:
        courses = MAJORS[self.major].get("Junior", [])
        final_courses = []
        for c in courses:
            if "AP Biology" in c and self.f_sci == "AP Biology":
                final_courses.append("AP Chemistry")
            else:
                final_courses.append(c)
        return final_courses

    @rx.var(cache=True)
    def senior_major_courses(self) -> list[str]:
        courses = MAJORS[self.major].get("Senior", [])
        courses = [c for c in courses if c != "AP Calculus BC or Multi-Variable Calculus"]

        # 🎯 LIU Pharmacy: AP Biology freshman year bumps the senior AP Biology
        # slot to AP Physics I, or AP Physics C: Mechanics if they also took
        # AP Physics I junior year.
        if self.major == "LIU Pharmacy" and self.f_sci == "AP Biology":
            took_ap_physics_jr = "AP Physics I" in self.junior_electives or self.j_phys == "AP Physics I"
            replacement = "AP Physics C: Mechanics" if took_ap_physics_jr else "AP Physics I"
            courses = [replacement if c == "AP Biology" else c for c in courses]

        return courses
    # --- Math Major Overrides and Filtering Backend ---
    @property
    def _math_keywords(self) -> list[str]:
        return ["Calculus", "Algebra II", "Geometry", "Pre-Calculus"]

    @rx.var
    def junior_math_override(self) -> str:
        if self.major == "Applied Mathematics":
            return "Mathematics Major Override"
        for c in self.junior_major_courses:
            if any(kw in c for kw in self._math_keywords):
                return c
            
        return ""

    @rx.var
    def senior_math_override(self) -> str:
        if self.major == "Applied Mathematics" or self.major == "LIU Pharmacy":
            return "Mathematics Major Override"
        for c in self.senior_major_courses:
            if any(kw in c for kw in self._math_keywords):
                return c
        return ""

    @rx.var
    def filtered_junior_major_courses(self) -> list[str]:
        override = self.junior_math_override
        return [c for c in self.junior_major_courses if c != override]

    @rx.var
    def filtered_senior_major_courses(self) -> list[str]:
        override = self.senior_math_override
        return [c for c in self.senior_major_courses if c != override]

    
    def _courses_for_year(self, year: str) -> set:
        if year == "Freshman":
            return set([
                "Design & Drafting for Production (DDP)", "English 9",
                self.f_hist, self.f_sci, self.f_lang, "Gym", "Lunch" 
            ] + self.freshman_math_courses + self.freshman_electives)

        if year == "Sophomore":
            math = "Pre-Calculus" if self.math_is_accelerated else "Algebra II"
            lang = (self.language_base + " III") if self.lang_is_accelerated else (self.language_base + " II")
            courses = ([
                "AP Computer Science Principles", "Health Education", self.so_eng, self.so_hist,
                "Regents Chemistry", math, lang, "Gym", "Lunch" 
            ] + self.sophomore_electives)
            
            return set(courses)
            
        if year == "Junior":
            math = self.j_math
            # Start with dropdowns and electives as a clean list
            courses = [ self.j_his, "Gym", "Lunch"] + self.junior_electives

            if self.major != "LIU Pharmacy":
                courses.append(self.j_eng)

            # 🎯 If they chose Economics in Junior year, dynamically inject Government into Junior year active tracker
            if (self.j_his == "Economics" or (
                "AP Macroeconomics / AP Microeconomics (One Semester each)" in self.junior_major_courses or 
                "AP Macroeconomics / AP Microeconomics (One Semester each)" in self.junior_electives
            )) and not self.has_econ_gov_conflict and not self.chose_government_jr and not self.has_gov_conflict:
                courses.append(self.sn_gov)
            
            if self.junior_math_override != "":
                courses.extend(self.filtered_junior_major_courses)
                courses.append(math)
            else:
                courses.extend(self.junior_major_courses)
                courses.append(math)
                
            if not self.lang_is_accelerated:
                courses.append(self.language_base + " III")
            if self.need_regents_physics_junior:
                courses.append(self.j_phys)
            return set(c for c in courses if c not in self.none_hisory_placeholders)
        
        if year == "Senior":

            math = self.sn_math if self.sn_math else "AP Calculus BC"
            courses = [self.sn_eng, "Gym", "Lunch"] + self.senior_electives
            
            if self.major == "Physics":
                if math and self.sn_math_active:
                    courses.append(math)
                courses.extend(list(self.selected_senior_physics))
                courses.append("Modern Physics") 
            else:
                if self.senior_math_override != "":
                    courses.extend(self.filtered_senior_major_courses)
                else:
                    courses.extend(self.senior_major_courses)
                if math and self.sn_math_active:
                    courses.append(math)
            # 🎯 Only track Senior History rules if they DID NOT choose Economics early in Junior year
            if self.has_macro:
                pass
            elif not self.chose_economics_jr:
                if self.chose_government_jr:
                    courses.append("Economics")
                else:
                    courses.append(self.sn_his)
                    if self.sn_his in ("AP US Government", "Participation in Government"):
                        courses.append(self.sn_ec)
                    if self.sn_his == "Economics" and not self.has_gov_conflict:
                        courses.append(self.sn_gov)

            return set(c for c in courses if c not in self.none_hisory_placeholders)
    @rx.var
    def all_courses(self) -> set[str]:
        return self._courses_for_year("Freshman") | self._courses_for_year("Sophomore") | self._courses_for_year("Junior") | self._courses_for_year("Senior") 
    
    def _evict_duplicates(self):
        fr = self._courses_for_year("Freshman")
        so = self._courses_for_year("Sophomore")
        jr = self._courses_for_year("Junior")

        prior_after_fr = fr
        prior_after_so = fr | so
        prior_after_jr = fr | so | jr

        if self.so_eng in prior_after_fr and self.so_eng != "English 10":
            self.so_eng = "English 10"
        if self.so_hist in prior_after_fr:
            self.so_hist = "Regents Global History"
        self.sophomore_electives = [c for c in self.sophomore_electives if c not in prior_after_fr or COURSE_CATALOG.get(c, {}).get("allow_multiple", False)]

        if self.j_eng in prior_after_so:
            self.j_eng = "English 11"
        if self.j_his in prior_after_so:
            self.j_his = "Regents United States History"
        if self.j_his == "Economics" and self.has_econ_gov_conflict:
            self.j_his = "Regents United States History"
            self.sn_his = "Economics"
        self.junior_electives = [c for c in self.junior_electives if c not in prior_after_so or COURSE_CATALOG.get(c, {}).get("allow_multiple", False)]

        if self.sn_eng in prior_after_jr:
            self.sn_eng = "Creative Writing"
        if self.sn_his in prior_after_jr:
            self.sn_his = "Participation in Government"
        if self.sn_ec in prior_after_jr:
            self.sn_his = "Participation in Government"
        if self.sn_gov in prior_after_jr:
            self.sn_gov = "Participation in Government"
        self.senior_electives = [c for c in self.senior_electives if c not in prior_after_jr or COURSE_CATALOG.get(c, {}).get("allow_multiple", False)]

    def _enforce_ap_limits(self):
        if self.junior_ap_count > 3:
            if "AP" in self.j_eng: self.j_eng = "English 11"
            if "AP" in self.j_his: self.j_his = "Regents United States History"
            if "AP" in self.j_phys: self.j_phys = "Regents Physics"
            if "AP" in self.j_math and not self.math_is_accelerated: self.j_math = "Pre-Calculus"
            while self.junior_ap_count > 3:
                ap_els = [c for c in self.junior_electives if "AP" in c]
                if not ap_els: break
                self.junior_electives.remove(ap_els[0])
        if self.senior_ap_count > 3:
            if "AP" in self.sn_eng: self.sn_eng = "Creative Writing"
            if "AP" in self.sn_his: self.sn_his = "Economics"
            if "AP" in self.sn_ec: self.sn_ec = "Economics"
            while self.senior_ap_count > 3:
                ap_els = [c for c in self.senior_electives if "AP" in c]
                if not ap_els: break
                self.senior_electives.remove(ap_els[0])

    def _enforce_period_limits(self) -> bool:
        
        year_to_electives = {
            "Freshman": "freshman_electives",
            "Sophomore": "sophomore_electives",
            "Junior": "junior_electives",
            "Senior": "senior_electives",
        }
        ok = True
        for year, attr in year_to_electives.items():
            while self._calculate_total_periods(year) > 10:
                electives = getattr(self, attr)
                if not electives:
                    ok = False
                    break  # over the cap with no electives left to cut — nothing more to auto-fix
                electives.pop(0) 
        return ok

    
    def compute_filtered_catalog(self) -> list[dict]:
        query = self.search_query.lower().strip()
        results = []
        
        active_courses = set()
        for yr in ["Freshman", "Sophomore", "Junior", "Senior"]:
            active_courses.update(self._courses_for_year(yr))

        current_grade = str(self.current_modal_grade).strip()

        # Headroom AND AP limit verification routing based on active modal view
        if current_grade == "Freshman":
            current_periods = self.freshman_periods
            ap_limit_reached = self.freshman_ap_count >= 3
        elif current_grade == "Sophomore":
            current_periods = self.sophomore_periods
            ap_limit_reached = self.sophomore_ap_count >= 3
        elif current_grade == "Junior":
            current_periods = self.junior_periods
            ap_limit_reached = self.junior_ap_count >= 3
        else:
            current_periods = self.senior_periods
            ap_limit_reached = self.senior_ap_count >= 3

        available_headroom = 10 - current_periods

        for name, data in COURSE_CATALOG.items():
            
            if not data.get("is_elective", False) or (name in active_courses and not data.get("allow_multiple", False)):
                continue
            
            if name in ["AP Macroeconomics with Government", "AP Microeconomics with Government"]:
                if not self.chose_economics_jr:
                    continue

            if name == "Linear Algebra" and not self.took_ap_calc_bc_junior:
                continue
            if name == "Multivariable Calculus" and not self.took_ap_calc_bc_junior:
                continue
            if name == "AP Physics C: Mechanics":
                if "AP Physics C: Electricity & Magnetism" in self._courses_for_year("Senior"):
                    continue

            if name == "AP Physics C: Electricity & Magnetism" and self.major == "Physics":  
                continue
            
            if current_grade == "Junior" and self.language_base == "Spanish":
                if name == "AP Spanish Language & Culture" and "Advanced Spanish" not in self._courses_for_year("Sophomore"):
                    continue

            if current_grade == "Junior" and self.language_base == "Chinese" and not self.lang_is_accelerated:
                if name == "AP Chinese Language & Culture":
                    continue
            if current_grade == "Junior" and self.language_base == "German" and not self.lang_is_accelerated:
                if name == "AP German Language & Culture":
                    continue

            if current_grade == "Junior" and self.language_base == "Italian" and not self.lang_is_accelerated:
                if name == "AP Italian Language & Culture":
                    continue
            if current_grade == "Junior" and self.language_base == "French" and not self.lang_is_accelerated:
                if name == "AP French Language & Culture":
                    continue

            if name in self.ap_lotes and self.language_base not in name:
                continue

            if name == "Advanced Spanish" and current_grade == "Junior":
                if "Spanish III" not in self._courses_for_year("Sophomore"):
                    continue

            if name == "Advanced Spanish" and current_grade == "Senior":
                if "Spanish III" not in self._courses_for_year("Junior"):
                    continue
            if name == "AP Spanish Language & Culture" and "Advanced Spanish" not in self.all_courses:
                continue
            

            if "available_grades" in data:
                allowed = data["available_grades"]
                if current_grade not in allowed:
                    continue
                
            else:
                if current_grade in ["Freshman", "Sophomore"]:
                    continue

            # ◄ RULE: Unless it explicitly says 2 periods, default it to 1 period
            raw_periods = data.get("periods", 1)
            semesters = data.get("semesters", 2)
            
            
            is_too_large = raw_periods > available_headroom
            is_ap_course = "AP" in name
            should_disable = is_too_large or (is_ap_course and ap_limit_reached)

            if query in name.lower() or query in data.get("description", "").lower():
                
                credits_formatted = f"{float(data.get('credits', 1.0)):.1f}"

                results.append({
                    "name": name, 
                    "description": data.get("description", ""),
                    "periods": raw_periods, # ◄ FIX: This now passes '1' or '2' to the badge display
                    "semesters" : semesters,
                    "credits": credits_formatted,
                    "is_elective" : data.get("is_elective", False),
                    "disabled": should_disable,
                    "grade_badges": _grade_badges(
                        data.get("is_elective", False),
                        data.get("grades", []),
                        data.get("available_grades", [])
                    ),
                })
        return results
   #login 

    @rx.var
    def filtered_catalog(self) -> list[dict]:
        return self.compute_filtered_catalog()
    
    @rx.var
    def filtered_catalog_count(self) -> int:
        return len(self.compute_filtered_catalog())
    @rx.var
    def dropdown_colors(self) -> dict[str, str]:
        return {name: data.get("color", "#1e2025") for name, data in COURSE_CATALOG.items()}


def _expand_taken(taken: set) -> set:
    expanded = set()
    for t in taken:
        expanded.add(t)
        if " / " in t:
            for part in t.split(" / "):
                clean = part.strip()
                expanded.add(clean)
                no_paren = re.sub(r'\s*\(.*?\)\s*$', '', clean).strip()
                if no_paren != clean:
                    expanded.add(no_paren)
    return expanded

def _filter(opts: list, current: str, ap_count: int, taken: set) -> list:
    expanded = _expand_taken(taken)
    other_ap_count = ap_count - (1 if "AP" in current else 0)
    at_limit = other_ap_count >= 3
    return [
        o for o in opts
        if (o == current or o not in expanded)
        and (o == current or "AP" not in o or not at_limit)
    ]

class DropdownOptions(PlannerState):
    @rx.var
    def f_hist_opts(self) -> list[str]:
        opts = ["Pre-AP World History", "AP Human Geography"]
        return _filter(opts, self.f_hist, self.freshman_ap_count, set())
    @rx.var
    def f_sci_opts(self) -> list[str]:
        has_bio = self.pre_credits.get("Biology/Living Environment", False)
        
        if not has_bio:
            opts = ["Regents Biology"]
        else:
            opts = ["Regents Physics", "AP Biology"]
        return _filter(opts, self.f_sci, self.freshman_ap_count , set())

    @rx.var
    def so_eng_opts(self) -> list[str]:
        opts = ["English 10", "AP Capstone Seminar"]
        return _filter(opts, self.so_eng, self.sophomore_ap_count , self._courses_for_year("Freshman"))
    @rx.var
    def so_hist_opts(self) -> list[str]:
        opts = ["AP European History", "AP World History", "Regents Global History"]
        return _filter(opts, self.so_hist, self.sophomore_ap_count , self._courses_for_year("Freshman"))
    @rx.var
    def j_eng_opts(self) -> list[str]:
        opts = ["AP English Language & Composition", "English 11"]
        prior = self._courses_for_year("Freshman") | self._courses_for_year("Sophomore")
        return _filter(opts, self.j_eng, self.junior_ap_count , prior)
    @rx.var
    def j_his_opts(self) -> list[str]:
        opts = ["AP United States History", "Regents United States History", "Economics", "AP African American Studies", "None (free period, econ senior year)"]
        if self.has_econ_gov_conflict:
            opts = [o for o in opts if o not in ("Economics",)]
        if not self.pre_credits.get("US History", False):
            opts = [o for o in opts if o not in ( "Economics", "None (free period, econ senior year)")]
        prior = self._courses_for_year("Freshman") | self._courses_for_year("Sophomore")
        return _filter(opts, self.j_his, self.junior_ap_count ,prior)
    @rx.var
    def j_math_opts(self) -> list[str]:
        opts = ["AP Calculus AB", "AP Calculus BC"] if self.math_is_accelerated else ["Pre-Calculus", "AP Calculus AB", "AP Calculus BC"]
        prior = self._courses_for_year("Freshman") | self._courses_for_year("Sophomore")
        return _filter(opts, self.j_math, self.junior_ap_count, prior)
    @rx.var
    def j_phys_opts(self) -> list[str]:
        return ["Regents Physics", "AP Physics I"]
    @rx.var
    def sn_calcs(self) -> list[str]:
        return ["AP Calculus BC", "AP Calculus AB"]
    @rx.var
    def sn_math_opts(self) -> list[str]:
        if self.major == "LIU Pharmacy":
            opts = ["Multivariable Calculus", "Linear Algebra", "AP Calculus BC"] if self.j_math == "AP Calculus AB" else ["Multivariable Calculus", "Linear Algebra"]
        elif self.senior_math_override:
            opts = ["Multivariable Calculus", "AP Calculus BC"] if self.j_math == "AP Calculus AB" else ["Multivariable Calculus"]    
        else:
            opts = ["Multivariable Calculus", "Linear Algebra", "AP Calculus BC"] if self.j_math == "AP Calculus AB" else ["Multivariable Calculus", "Linear Algebra"]
        prior = self._courses_for_year("Freshman") | self._courses_for_year("Sophomore") | self._courses_for_year("Junior")
        return _filter(opts, self.sn_math, False, prior)
    @rx.var
    def sn_eng_opts(self) -> list[str]:
        opts = ["AP Capstone Research", "AP English Literature & Composition", "Creative Writing", "Film & Literature", "Graphic Novels", "Journalism", "Science Fiction & Fantasy", "Mystery, Horror & the Supernatural", "Live, Love, & Death", "Drama", "Yearbook"] if self.so_eng == "AP Capstone Seminar" else [ "AP English Literature & Composition", "Creative Writing", "Film & Literature", "Graphic Novels", "Journalism", "Science Fiction & Fantasy", "Mystery, Horror & the Supernatural", "Live, Love, & Death", "Drama", "Yearbook"]
        prior = self._courses_for_year("Freshman") | self._courses_for_year("Sophomore") | self._courses_for_year("Junior")
        return _filter(opts, self.sn_eng, self.senior_ap_count , prior)
    @rx.var
    def sn_his_opts(self) -> list[str]:
        if self.has_econ_gov_conflict:
            opts = ["AP US Government", "Participation in Government"]
        else:
            opts = ["AP Macroeconomics with Government", "AP Microeconomics with Government", "Economics", "AP US Government", "Participation in Government"]
            if self.pre_credits.get("US History", False):
                opts.append("None (free period, econ junior year)")
        if self.chose_government_jr:
            opts = [o for o in opts if o != "Participation in Government"]

        prior = self._courses_for_year("Freshman") | self._courses_for_year("Sophomore") | self._courses_for_year("Junior")
        return _filter(opts, self.sn_his, self.senior_ap_count, prior)
      
    @rx.var
    def sn_ec_opts(self) -> list[str]:
        opts = [#"AP Macroeconomics with Government", 
                #"AP Microeconomics with Government", 
                "Economics"]
        prior = self._courses_for_year("Freshman") | self._courses_for_year("Sophomore") | self._courses_for_year("Junior")
        return _filter(opts, self.sn_ec, self.senior_ap_count , prior)
    @rx.var
    def sn_gov_opts(self) -> list[str]:
        opts = ["AP US Government", "Participation in Government"]
        prior = self._courses_for_year("Freshman") | self._courses_for_year("Sophomore") | self._courses_for_year("Junior")
        return _filter(opts, self.sn_gov, self.senior_ap_count, prior)



MAJOR_COURSE_COLOR = req_courses