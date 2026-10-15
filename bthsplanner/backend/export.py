import io
import openpyxl

YEAR_COLS = {"Freshman": 2, "Sophomore": 3, "Junior": 4, "Senior": 5}
TOTAL_COL = 6


def build_plan_xlsx(
    major: str,
    year_courses: dict[str, list[str]],
    course_counts: dict[str, int],
    periods: dict[str, int],
    ap_counts: dict[str, int],
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "4-Year Plan"

    ws.cell(row=1, column=1, value="Major")
    ws.cell(row=1, column=2, value=major)

    ws.cell(row=3, column=1, value="Year")
    for year, col in YEAR_COLS.items():
        ws.cell(row=3, column=col, value=year)

    # courses stacked under each year
    for year, col in YEAR_COLS.items():
        for i, course in enumerate(year_courses[year]):
            ws.cell(row=4 + i, column=col, value=course)

    # summary rows, one per stat, with a total in the last column
    max_courses = max(len(c) for c in year_courses.values())
    summary_start = 4 + max_courses + 1
    summaries = [
        ("Courses", course_counts),
        ("Periods", periods),
        ("AP Count", ap_counts),
    ]
    for offset, (label, counts) in enumerate(summaries):
        row = summary_start + offset
        ws.cell(row=row, column=1, value=label)
        for year, col in YEAR_COLS.items():
            ws.cell(row=row, column=col, value=counts[year])
        ws.cell(row=row, column=TOTAL_COL, value=sum(counts.values()))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()